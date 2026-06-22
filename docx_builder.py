"""Maqola dict'idan OAK uslubidagi Word (.docx) hujjat tuzish."""
from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from chart_builder import render_chart

_LABELS = {
    "annotation": {"uz": "Annotatsiya", "ru": "Аннотация", "en": "Abstract"},
    "keywords": {
        "uz": "Kalit so'zlar",
        "ru": "Ключевые слова",
        "en": "Keywords",
    },
    "table": {"uz": "jadval", "ru": "таблица"},
    "figure": {"uz": "rasm", "ru": "рисунок"},
}

_SECTION_TITLES = {
    "uz": {
        "introduction": "KIRISH",
        "main_part": "ASOSIY QISM",
        "results": "NATIJALAR VA ULARNING TAHLILI",
        "conclusion": "XULOSA",
        "references": "FOYDALANILGAN ADABIYOTLAR RO'YXATI",
    },
    "ru": {
        "introduction": "ВВЕДЕНИЕ",
        "main_part": "ОСНОВНАЯ ЧАСТЬ",
        "results": "РЕЗУЛЬТАТЫ И ИХ ОБСУЖДЕНИЕ",
        "conclusion": "ЗАКЛЮЧЕНИЕ",
        "references": "СПИСОК ИСПОЛЬЗОВАННОЙ ЛИТЕРАТУРЫ",
    },
}


def _add_body(doc: Document, text: str) -> None:
    """Ko'p paragrafli matnni qo'shadi (bo'sh qatorlar bo'yicha bo'linadi)."""
    for chunk in (text or "").split("\n\n"):
        chunk = chunk.strip()
        if not chunk:
            continue
        p = doc.add_paragraph(chunk)
        p.paragraph_format.first_line_indent = Pt(18)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _add_tables(doc: Document, tables: list, lang: str) -> None:
    """Premium jadvallarni qo'shadi (sarlavha + jadval)."""
    for idx, tbl in enumerate(tables, 1):
        headers = [str(h) for h in (tbl.get("headers") or [])]
        rows = tbl.get("rows") or []
        if not headers and not rows:
            continue

        cap = doc.add_paragraph()
        label = _LABELS["table"].get(lang, _LABELS["table"]["uz"])
        cap.add_run(f"{idx}-{label}. {tbl.get('title', '')}".strip()).bold = True

        ncols = max(len(headers), max((len(r) for r in rows), default=0))
        if ncols == 0:
            continue
        table = doc.add_table(rows=0, cols=ncols)
        table.style = "Table Grid"

        if headers:
            hcells = table.add_row().cells
            for i in range(ncols):
                text = headers[i] if i < len(headers) else ""
                hcells[i].text = str(text)
                for p in hcells[i].paragraphs:
                    for run in p.runs:
                        run.bold = True
        for r in rows:
            cells = table.add_row().cells
            for i in range(ncols):
                cells[i].text = str(r[i]) if i < len(r) else ""
        doc.add_paragraph()


def _add_charts(doc: Document, charts: list, lang: str) -> None:
    """Premium diagrammalarni (matplotlib rasm) qo'shadi."""
    for idx, chart in enumerate(charts, 1):
        png = render_chart(chart)
        if png is None:
            continue
        doc.add_picture(png, width=Inches(5.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label = _LABELS["figure"].get(lang, _LABELS["figure"]["uz"])
        cap.add_run(f"{idx}-{label}. {chart.get('title', '')}".strip()).italic = True
        doc.add_paragraph()


def build_docx(article: dict, author: str, lang: str) -> io.BytesIO:
    """Maqoladan .docx tuzib, BytesIO oqimini qaytaradi."""
    lang = lang if lang in _SECTION_TITLES else "uz"
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)

    # UDK
    udk = doc.add_paragraph()
    udk.add_run(f"UDK: {article.get('udk', '')}").bold = True

    # Sarlavha (interfeys tilida)
    title = article.get("title", {})
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = h.add_run(title.get(lang) or title.get("uz", ""))
    run.bold = True
    run.font.size = Pt(15)

    # Muallif
    if author and author.strip() not in {"—", "-"}:
        a = doc.add_paragraph()
        a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        a.add_run(author).italic = True

    doc.add_paragraph()

    # Annotatsiya va kalit so'zlar — uchta tilda
    annotation = article.get("annotation", {})
    keywords = article.get("keywords", {})
    for code in ("uz", "ru", "en"):
        ann = annotation.get(code)
        if ann:
            p = doc.add_paragraph()
            p.add_run(f"{_LABELS['annotation'][code]}. ").bold = True
            p.add_run(ann)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        kw = keywords.get(code) or []
        if kw:
            p = doc.add_paragraph()
            p.add_run(f"{_LABELS['keywords'][code]}: ").bold = True
            p.add_run(", ".join(kw))
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_paragraph()

    # Asosiy bo'limlar
    titles = _SECTION_TITLES[lang]
    for key in ("introduction", "main_part", "results", "conclusion"):
        heading = doc.add_paragraph()
        heading.add_run(titles[key]).bold = True
        _add_body(doc, article.get(key, ""))
        # Jadval/diagrammalar natijalardan keyin joylashtiriladi (premium)
        if key == "results":
            _add_tables(doc, article.get("tables") or [], lang)
            _add_charts(doc, article.get("charts") or [], lang)

    # Adabiyotlar
    refs = article.get("references") or []
    if refs:
        heading = doc.add_paragraph()
        heading.add_run(titles["references"]).bold = True
        for i, ref in enumerate(refs, 1):
            doc.add_paragraph(f"{i}. {ref}")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
