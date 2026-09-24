"""Generate synthetic sample documents (uz / ru / en) as DOCX, PDF, TXT, plus a scanned PDF.

    python scripts/make_samples.py [output_dir]

All content comes from tests/fixtures/sample_texts.py and is synthetic.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from tests.fixtures.sample_texts import SAMPLES, dissertation_outline  # noqa: E402

FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "/Library/Fonts/Arial Unicode.ttf",
]


def font_path() -> str | None:
    return next((f for f in FONT_CANDIDATES if Path(f).exists()), None)


def make_docx(lang: str) -> bytes:
    import docx
    from docx.enum.text import WD_BREAK

    d = docx.Document()
    first_chapter_seen = False
    for kind, text in dissertation_outline(lang):
        if kind == "title":
            d.add_heading(text, level=0)
        elif kind == "heading":
            is_sub = text[:4].strip()[:1].isdigit() and "." in text[:4]
            if not is_sub:
                # new top-level part starts on a new page (real page breaks)
                p = d.add_paragraph()
                p.add_run().add_break(WD_BREAK.PAGE)
                first_chapter_seen = True
            d.add_heading(text, level=2 if is_sub else 1)
        else:
            d.add_paragraph(text)
    _ = first_chapter_seen
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Guruh / Group", "%"
    t.cell(1, 0).text, t.cell(1, 1).text = "A", "78"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_txt(lang: str) -> bytes:
    parts = []
    for kind, text in dissertation_outline(lang):
        parts.append(text)
    return ("\n\n".join(parts) + "\n").encode("utf-8")


def make_pdf(lang: str) -> bytes:
    import pymupdf as fitz

    fp = font_path()
    pdf = fitz.open()
    page = None
    y = 0.0
    W, H, M = 595, 842, 60

    def new_page():
        nonlocal page, y
        page = pdf.new_page(width=W, height=H)
        if fp:
            page.insert_font(fontname="dv", fontfile=fp)
        y = M
        page.insert_text((W / 2 - 5, H - 30), str(pdf.page_count), fontname="dv" if fp else "helv", fontsize=9)

    new_page()
    for kind, text in dissertation_outline(lang):
        size = 15 if kind == "title" else 12.5 if kind == "heading" else 10.5
        if kind == "heading" and not text[:1].isdigit():
            new_page()
        rect = fitz.Rect(M, y, W - M, H - M)
        rc = page.insert_textbox(rect, text, fontname="dv" if fp else "helv", fontsize=size)
        if rc < 0:  # does not fit: next page
            new_page()
            rect = fitz.Rect(M, y, W - M, H - M)
            rc = page.insert_textbox(rect, text, fontname="dv" if fp else "helv", fontsize=size)
        used = (H - M - y) - max(rc, 0)
        y += used + size * 1.2
    pdf.subset_fonts()
    return pdf.tobytes(garbage=4, deflate=True)


def make_scanned_pdf(lang: str) -> bytes:
    """Rasterise the text PDF's first pages into images (no text layer) to exercise OCR."""
    import pymupdf as fitz

    src = fitz.open(stream=make_pdf(lang), filetype="pdf")
    out = fitz.open()
    for pno in range(min(2, src.page_count)):
        pix = src[pno].get_pixmap(dpi=150, colorspace=fitz.csGRAY)
        page = out.new_page(width=src[pno].rect.width, height=src[pno].rect.height)
        page.insert_image(page.rect, stream=pix.tobytes("jpg", jpg_quality=70))
    return out.tobytes(garbage=4, deflate=True)


def main(out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    for lang in SAMPLES:
        (out_dir / f"sample_{lang}.docx").write_bytes(make_docx(lang))
        (out_dir / f"sample_{lang}.pdf").write_bytes(make_pdf(lang))
        (out_dir / f"sample_{lang}.txt").write_bytes(make_txt(lang))
    (out_dir / "sample_en_scanned.pdf").write_bytes(make_scanned_pdf("en"))
    print(f"samples written to {out_dir}")


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "tests" / "fixtures" / "samples")
