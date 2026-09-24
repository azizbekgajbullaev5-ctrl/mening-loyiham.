"""A ~300-page synthetic dissertation must be processed in chunks, quickly, with progress tracking."""
import io
import time

import docx
import pytest

from app.core.database import SessionLocal
from app.models import Analysis
from tests.conftest import upload
from tests.fixtures.sample_texts import SAMPLES


def _big_docx(chapters: int = 5, sections: int = 6, paras: int = 45) -> bytes:
    s = SAMPLES["uz"]
    pool = s["human"] + s["ai"]
    d = docx.Document()
    d.add_heading("KATTA DISSERTATSIYA (sintetik)", 0)
    d.add_heading("KIRISH", 1)
    d.add_paragraph(s["human"][0])
    roman = ["I", "II", "III", "IV", "V", "VI"]
    k = 0
    for c in range(chapters):
        d.add_heading(f"{roman[c]} BOB. BOB NOMI {c + 1}", 1)
        for sec in range(sections):
            d.add_heading(f"{c + 1}.{sec + 1}. Bo'lim nomi {sec + 1}", 2)
            for _ in range(paras):
                # vary numbers so passages are not all identical
                d.add_paragraph(pool[k % len(pool)].replace("2022", str(1990 + k % 30)) + f" [{k % 40 + 1}]")
                k += 1
    d.add_heading("XULOSA", 1)
    d.add_paragraph(s["ai"][3])
    d.add_heading("FOYDALANILGAN ADABIYOTLAR RO'YXATI", 1)
    for i in range(40):
        d.add_paragraph(f"{i + 1}. Muallif {i}. Asar nomi. – Toshkent, {1990 + i % 30}. (namuna)")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.mark.slow
def test_300_page_dissertation(user_client):
    data = _big_docx()
    t0 = time.monotonic()
    item = upload(user_client, "big.docx", data, depth="standard")
    elapsed = time.monotonic() - t0
    a = user_client.get(f"/api/analyses/{item['analysis']['id']}").json()
    assert a["status"] == "completed", a["error"]
    assert a["version"]["word_count"] > 75_000  # ~300 pages at ~280 words/page
    assert a["version"]["page_count"] >= 250
    res = a["result"]
    assert res["ai"]["passages_analyzed"] > 300
    sections = user_client.get(f"/api/analyses/{a['id']}/sections").json()
    assert sum(1 for s in sections if s["kind"] == "chapter") == 5
    assert sum(1 for s in sections if s["kind"] == "section") == 30
    assert elapsed < 180, f"too slow: {elapsed:.1f}s"
    # heavy re-use inside the synthetic text must be reported by the similarity module
    assert res["similarity"]["internal"] > 50
    # only flagged passages are stored with text (privacy)
    with SessionLocal() as db:
        stored = db.get(Analysis, a["id"])
        assert len(stored.passages) == res["ai"]["passages_flagged"] < res["ai"]["passages_analyzed"]
    r = user_client.get(f"/api/analyses/{a['id']}/report", params={"format": "pdf"})
    assert r.status_code == 200 and len(r.content) > 10_000
