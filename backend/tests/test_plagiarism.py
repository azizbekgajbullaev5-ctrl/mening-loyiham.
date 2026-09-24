"""Plagiarism system: corpus, harvesting, web check, algorithm, report, tricks."""
import io
import json

import docx
import httpx
import pymupdf
import pytest
from docx.shared import Pt, RGBColor
from fastapi.testclient import TestClient
from pydantic import SecretStr

from app.core import config
from app.core.database import SessionLocal
from app.main import app
from app.models import CorpusJobItem, RefDocument, RefFingerprint, RefVector, WebPageCache
from app.plagiarism import integrity, web
from app.plagiarism.textnorm import canonical_words
from tests.conftest import register, upload
from tests.fixtures.sample_texts import SAMPLES

UZ = SAMPLES["uz"]
PARAPHRASE = (
    "Olingan natijalar biz kutgandan birmuncha farqli bo'ldi. Platformada ishlagan talabalar o'rtacha 3,9 ball to'plagan bo'lsa, "
    "an'anaviy guruhlarda bu ko'rsatkich 3,7 ballni tashkil etdi; ushbu farq statistik ahamiyatga ega emas (p = 0,12). Biroq topshiriqlarni "
    "o'z vaqtida topshirish darajasi keskin farqlandi: 41 foizga nisbatan 78 foiz. Boshqacha aytganda, platforma bilim darajasini emas, "
    "intizomni o'zgartirdi. Suhbat chog'ida talabalar eng foydali jihat sifatida eslatmalar va muddatlarning ko'rinib turishini qayd etishdi."
)


def to_cyrillic(text: str) -> str:
    m = {"o'": "ў", "g'": "ғ", "sh": "ш", "ch": "ч", "a": "а", "b": "б", "d": "д", "e": "е", "f": "ф", "g": "г", "h": "ҳ", "i": "и",
         "j": "ж", "k": "к", "l": "л", "m": "м", "n": "н", "o": "о", "p": "п", "q": "қ", "r": "р", "s": "с", "t": "т", "u": "у",
         "v": "в", "x": "х", "y": "й", "z": "з"}
    out = text.lower()
    for k in ("o'", "g'", "sh", "ch"):
        out = out.replace(k, m[k])
    return "".join(m.get(c, c) for c in out)


def make_docx(paragraphs: list[str], headings: dict[int, str] | None = None) -> bytes:
    d = docx.Document()
    for i, p in enumerate(paragraphs):
        if headings and i in headings:
            d.add_heading(headings[i], level=1)
        d.add_paragraph(p)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


@pytest.fixture
def admin():
    """The first user of a fresh DB is admin; in the shared test DB we promote explicitly."""
    from app.models import User

    with TestClient(app) as c:
        info = register(c)
        with SessionLocal() as db:
            db.get(User, info["user"]["id"]).is_admin = True
            db.commit()
        yield c


@pytest.fixture(autouse=True)
def _clean_corpus():
    with SessionLocal() as db:
        db.query(RefFingerprint).delete()
        db.query(RefVector).delete()
        db.query(RefDocument).delete()
        db.query(WebPageCache).delete()
        db.commit()
    from app.plagiarism.corpus import vector_index

    vector_index.invalidate()
    yield


def corpus_upload(client, files: list[tuple[str, bytes]], kind="article"):
    r = client.post("/api/corpus/upload", files=[("files", (n.rsplit("/", 1)[-1], d, "application/octet-stream")) for n, d in files],
                    data={"doc_kind": kind, "paths": [n for n, _ in files]})
    assert r.status_code == 201, r.text
    return r.json()


# ---------------------------------------------------------------- normalisation & tricks
def test_transliteration_and_homoglyphs():
    latin = canonical_words("O'zbekiston Respublikasi oliy ta'lim tizimi")
    assert canonical_words("Ўзбекистон Республикаси олий таълим тизими") == latin
    assert canonical_words("Oʻzbekiston Rеspublikаsi oliy ta’lim tizimi") == latin  # Cyrillic е/а inside Latin words, ’ apostrophe
    assert canonical_words("tizi​mi") == ["tizimi"]


def test_integrity_scan_detects_tricks():
    from app.document_processing.types import Block, ExtractedDocument

    blocks = [Block(0, "Bu mаtn kirill hаrflari bilan аlmashtirilgаn so'zlarni o'z ichigа olаdi."), Block(1, "Nor​mal​ so'z​lar " * 5),
              Block(2, "p l a g i a t va y a n a b i r i va t e x n i k a")]
    doc = ExtractedDocument("txt", blocks, 1, hidden_fragments=[{"page": 1, "reason": "white", "chars": 300, "sample": "junk"}])
    codes = {i["code"]: i for i in integrity.scan(doc)["items"]}
    assert codes["homoglyphs"]["count"] >= 5
    assert codes["invisible_chars"]["count"] >= 10
    assert codes["hidden_text"]["severity"] == "high"
    assert "spaced_letters" in codes


def test_hidden_white_text_is_excluded_and_reported(user_client):
    d = docx.Document()
    p = d.add_paragraph(UZ["human"][0] + " ")
    junk = p.add_run(" ".join(["originallikni oshiruvchi yashirin so'z"] * 30))
    junk.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    tiny = d.add_paragraph().add_run("mayda " * 40)
    tiny.font.size = Pt(1)
    b = io.BytesIO()
    d.save(b)
    aid = upload(user_client, "hidden.docx", b.getvalue(), depth="quick")["analysis"]["id"]
    pl = user_client.get(f"/api/analyses/{aid}/plagiarism").json()
    items = {i["code"]: i for i in pl["integrity"]["items"]}
    assert items["hidden_text"]["reasons"] == {"white": 1, "tiny": 1}
    content = user_client.get(f"/api/documents/{upload(user_client, 'h2.docx', b.getvalue(), depth='quick')['document']['id']}/content").json()
    assert "yashirin" not in " ".join(x["text"] for x in content["paragraphs"])


# ---------------------------------------------------------------- reference corpus
def test_corpus_bulk_upload_stats_dedupe_delete(admin):
    files = [("darsliklar/kitob1.docx", make_docx(UZ["human"])), ("maqolalar/maqola1.docx", make_docx(SAMPLES["ru"]["human"]))]
    res = corpus_upload(admin, files, kind="textbook")
    job = admin.get(f"/api/corpus/jobs/{res['job']['id']}").json()
    assert job["status"] == "completed" and job["added"] == 2 and job["failed"] == 0
    st = admin.get("/api/corpus/stats").json()
    assert st["documents"] == 2 and st["fingerprints"] > 50 and st["vectors"] > 0 and st["database_bytes"]
    listing = admin.get("/api/corpus/documents", params={"q": "darslik"}).json()
    assert listing["total"] == 1 and listing["items"][0]["folder"] == "darsliklar" and listing["items"][0]["doc_kind"] == "textbook"
    # originals are not kept after indexing
    with SessionLocal() as db:
        assert all(i.storage_key is None for i in db.query(CorpusJobItem).all())
    # the same file again is recognised as a duplicate
    again = corpus_upload(admin, files[:1])
    assert admin.get(f"/api/corpus/jobs/{again['job']['id']}").json()["skipped"] == 1
    rid = listing["items"][0]["id"]
    assert admin.delete(f"/api/corpus/documents/{rid}").status_code == 200
    assert admin.get("/api/corpus/stats").json()["documents"] == 1
    with SessionLocal() as db:
        assert db.query(RefFingerprint).filter_by(ref_doc_id=rid).count() == 0


def test_corpus_management_requires_admin(user_client):
    r = user_client.post("/api/corpus/upload", files=[("files", ("a.docx", make_docx(UZ["human"]), "application/octet-stream"))])
    assert r.status_code == 403
    assert user_client.post("/api/corpus/harvest", json={"sources": ["openalex"], "queries": ["x"]}).status_code == 403
    assert user_client.get("/api/corpus/stats").status_code == 200  # read-only is allowed


# ---------------------------------------------------------------- algorithm end-to-end
def test_originality_borrowing_citation(admin):
    corpus_upload(admin, [("src/manba.docx", make_docx(UZ["human"]))])
    checked = make_docx(
        [
            "Ushbu ishda men o'zimning kuzatuvlarimni bayon qilaman va ularni boshqa hech qayerdan olmaganman. " * 4,
            UZ["human"][1],  # verbatim copy (Latin)
            to_cyrillic(UZ["human"][2]),  # the same source copied in Cyrillic script
            "Muallif ta'kidlaganidek: «" + UZ["human"][3][:260] + "» [3].",  # properly cited quotation
            "E = mc^2 + 2x",  # formula (excluded)
            "1. Karimov A. Talabalar mustaqil ishini nazorat qilish. 2021.",
        ],
        headings={5: "FOYDALANILGAN ADABIYOTLAR"},
    )
    aid = upload(admin, "tekshir.docx", checked, doc_type="article")["analysis"]["id"]
    pl = admin.get(f"/api/analyses/{aid}/plagiarism").json()
    assert pl["borrowing"] > 35 and pl["citation"] > 5 and pl["originality"] > 10
    assert abs(pl["originality"] + pl["borrowing"] + pl["citation"] - 100) < 0.05
    top = pl["sources"][0]
    assert top["title"] and top["module"] == "corpus" and top["share_text"] >= pl["borrowing"] - 0.1
    assert pl["exclusions"].get("references", 0) > 0 and pl["exclusions"].get("formulas", 0) > 0
    classes = {s[4] for s in pl["spans"]}
    assert {"b", "c"} <= classes
    summary = admin.get(f"/api/analyses/{aid}").json()
    assert summary["originality"] == pl["originality"]


def test_paraphrase_is_detected(admin):
    corpus_upload(admin, [("src/p.docx", make_docx(UZ["human"]))])
    aid = upload(admin, "para.docx", make_docx(["Mening kirish so'zim bu yerda va u mutlaqo boshqa mavzuda yozilgan. " * 3, PARAPHRASE]))["analysis"]["id"]
    pl = admin.get(f"/api/analyses/{aid}/plagiarism").json()
    assert pl["paraphrase_share"] > 10, pl
    assert any(s[4] == "p" for s in pl["spans"])
    assert pl["sources"][0]["paraphrase_words"] > 0


def test_unrelated_text_is_original(admin):
    corpus_upload(admin, [("src/x.docx", make_docx(UZ["human"]))])
    aid = upload(admin, "orig.docx", make_docx(SAMPLES["en"]["human"]))["analysis"]["id"]
    pl = admin.get(f"/api/analyses/{aid}/plagiarism").json()
    assert pl["originality"] > 95 and not pl["sources"]


# ---------------------------------------------------------------- internet check
class FakeBrave:
    calls = 0

    def __init__(self, *a, **k):
        pass

    def search(self, q):
        FakeBrave.calls += 1
        return [{"url": "https://example.org/article", "title": "Onlayn maqola"}, {"url": "https://example.org/other", "title": "Boshqa"}]


class FakeFetcher:
    fetched: list = []

    def __init__(self, *a, **k):
        pass

    def fetch_bytes(self, url):
        FakeFetcher.fetched.append(url)
        if url.endswith("article"):
            return "text/html", f"<html><head><title>Onlayn maqola</title></head><body><p>{UZ['human'][2]}</p></body></html>".encode()
        return "text/html", b"<html><body>Unrelated page about cooking recipes and travel.</body></html>"

    def fetch_text(self, url):
        ctype, data = self.fetch_bytes(url)
        return web._to_text(data, ctype)


def test_web_check_estimate_confirm_and_cache(user_client, monkeypatch):
    monkeypatch.setattr(config.get_settings(), "BRAVE_API_KEY", SecretStr("test-key"))
    monkeypatch.setattr(web, "BraveClient", FakeBrave)
    monkeypatch.setattr(web, "Fetcher", FakeFetcher)
    FakeBrave.calls, FakeFetcher.fetched = 0, []
    data = make_docx(["O'zim yozgan kirish qismi, bu yerda hech qanday nusxa yo'q. " * 3, UZ["human"][2], UZ["human"][3]])
    # keep_for_similarity=false: otherwise the 2nd upload is matched against the 1st and needs no web queries at all
    item = upload(user_client, "web.docx", data, web_check="true", keep_for_similarity="false")
    aid = item["analysis"]["id"]
    a = user_client.get(f"/api/analyses/{aid}").json()
    assert a["status"] == "awaiting_confirmation"
    est = a["web_estimate"]
    assert est["queries"] >= 3 and est["cost_usd"] == pytest.approx(est["queries"] * 5.0 / 1000) and est["configured"]
    assert FakeBrave.calls == 0  # nothing is spent before confirmation
    r = user_client.post(f"/api/analyses/{aid}/confirm", json={"web_check": True})
    assert r.status_code == 200
    pl = user_client.get(f"/api/analyses/{aid}/plagiarism").json()
    web_src = [s for s in pl["sources"] if s["module"] == "web"]
    assert web_src and web_src[0]["url"] == "https://example.org/article" and web_src[0]["share_text"] > 20
    stats = pl["modules"]["web_stats"]
    assert 0 < stats["queries_used"] <= est["queries"] and stats["fetched_pages"] == 2
    assert all(not q.startswith('"') for q in stats["queries"])  # plain queries, not strict exact phrases
    pages = {p["url"]: p for p in stats["pages"]}
    assert pages["https://example.org/article"]["status"] == "ok" and pages["https://example.org/article"]["source_index"] is not None
    assert pages["https://example.org/other"]["source_index"] is None
    assert stats["query_log"] and stats["results_total"] > 0
    # second check of the same file: the earlier copy is not a source, so the web is still searched; pages are cached
    FakeFetcher.fetched = []
    aid2 = upload(user_client, "web2.docx", data, web_check="true")["analysis"]["id"]
    user_client.post(f"/api/analyses/{aid2}/confirm", json={"web_check": True})
    assert FakeFetcher.fetched == []
    pl2 = user_client.get(f"/api/analyses/{aid2}/plagiarism").json()
    assert pl2["modules"]["web_stats"]["cached_pages"] == 2
    assert pl2["modules"]["duplicates"][0]["excluded"] and not [s for s in pl2["sources"] if s["module"] == "own"]
    pdf = user_client.get(f"/api/analyses/{aid2}/report", params={"kind": "plagiarism"})
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    text = "".join(page.get_text() for page in pymupdf.open(stream=pdf.content, filetype="pdf"))
    assert "https://example.org/article" in text and "avval yuklangan" in text


def test_web_check_reports_why_nothing_was_compared(user_client, monkeypatch):
    class NoResults(FakeBrave):
        def search(self, q):
            return []

    monkeypatch.setattr(config.get_settings(), "BRAVE_API_KEY", SecretStr("test-key"))
    monkeypatch.setattr(web, "BraveClient", NoResults)
    aid = upload(user_client, "nores.docx", make_docx(UZ["human"]), web_check="true", keep_for_similarity="false")["analysis"]["id"]
    user_client.post(f"/api/analyses/{aid}/confirm", json={"web_check": True})
    stats = user_client.get(f"/api/analyses/{aid}/plagiarism").json()["modules"]["web_stats"]
    assert stats["queries_used"] > 0 and stats["results_total"] == 0 and stats["queries_without_results"] == stats["queries_used"]
    assert any(e.startswith("brave_no_results") for e in stats["errors"])


def test_web_page_failures_are_listed_with_reason(user_client, monkeypatch):
    class Failing(FakeFetcher):
        def fetch_bytes(self, url):
            if url.endswith("article"):
                raise httpx.ConnectTimeout("slow")
            request = httpx.Request("GET", url)
            raise httpx.HTTPStatusError("forbidden", request=request, response=httpx.Response(403, request=request))

    monkeypatch.setattr(config.get_settings(), "BRAVE_API_KEY", SecretStr("test-key"))
    monkeypatch.setattr(web, "BraveClient", FakeBrave)
    monkeypatch.setattr(web, "Fetcher", Failing)
    aid = upload(user_client, "fail.docx", make_docx(UZ["human"]), web_check="true", keep_for_similarity="false")["analysis"]["id"]
    user_client.post(f"/api/analyses/{aid}/confirm", json={"web_check": True})
    stats = user_client.get(f"/api/analyses/{aid}/plagiarism").json()["modules"]["web_stats"]
    assert {p["url"]: p["status"] for p in stats["pages"]} == {"https://example.org/article": "timeout", "https://example.org/other": "http_403"}
    assert stats["failed_pages"] == 2 and any(e.startswith("pages_failed") for e in stats["errors"])


def test_query_count_scales_with_document_size():
    assert web.planned_queries(29205) == 74  # 1 query per 400 words
    assert web.planned_queries(200000) == config.get_settings().WEB_MAX_QUERIES
    assert web.planned_queries(300) == 3


# ---------------------------------------------------------------- earlier copies of the same document
def _own_sources(client, aid):
    pl = client.get(f"/api/analyses/{aid}/plagiarism").json()
    return pl, [s for s in pl["sources"] if s["module"] == "own"]


def test_reuploaded_document_is_not_its_own_source(user_client):
    data = make_docx(UZ["human"])
    upload(user_client, "A.N.Xasanova o'quv qo'llanma.docx", data)
    aid = upload(user_client, "A.N.Xasanova o'quv qo'llanma.docx", data)["analysis"]["id"]
    pl, own = _own_sources(user_client, aid)
    assert not own and pl["originality"] > 95
    dup = pl["modules"]["duplicates"][0]
    assert dup["excluded"] and {"same_file", "same_name"} <= set(dup["reasons"])
    sim = user_client.get(f"/api/analyses/{aid}/similarity", params={"match_type": "cross_document"}).json()
    assert sim["matches"] == []  # the "O'xshashlik" tab does not compare it with its own copy either


def test_same_text_under_another_name_is_a_copy(user_client):
    upload(user_client, "qollanma_v1.docx", make_docx(UZ["human"]))
    aid = upload(user_client, "yangi nom.docx", make_docx(UZ["human"] + ["Qo'shimcha kichik jumla."]))["analysis"]["id"]
    pl, own = _own_sources(user_client, aid)
    assert not own and pl["modules"]["duplicates"][0]["reasons"] == ["same_text"]


def test_same_name_different_work_is_still_compared(user_client):
    upload(user_client, "dissertatsiya.docx", make_docx(UZ["human"][:3]))
    other = ["Bu boshqa talabaning ishi, u o'z kuzatuvlari haqida yozadi va natijalarni tahlil qiladi. " * 4, UZ["human"][1]]
    aid = upload(user_client, "dissertatsiya.docx", make_docx(other))["analysis"]["id"]
    pl, own = _own_sources(user_client, aid)
    assert own, "a different work with the same file name must still be compared"
    dup = pl["modules"]["duplicates"][0]
    assert dup["reasons"] == ["same_name"] and not dup["excluded"]


def test_web_check_can_be_declined(user_client, monkeypatch):
    monkeypatch.setattr(config.get_settings(), "BRAVE_API_KEY", SecretStr("test-key"))
    monkeypatch.setattr(web, "BraveClient", FakeBrave)
    FakeBrave.calls = 0
    aid = upload(user_client, "noweb.docx", make_docx(UZ["human"]), web_check="true")["analysis"]["id"]
    user_client.post(f"/api/analyses/{aid}/confirm", json={"web_check": False})
    a = user_client.get(f"/api/analyses/{aid}").json()
    assert a["status"] == "completed" and not a["web_check"] and FakeBrave.calls == 0


@pytest.mark.parametrize("url", ["http://127.0.0.1/x", "http://10.1.2.3/", "http://localhost:8000/", "file:///etc/passwd", "ftp://example.org/", "http://example.org:22/"])
def test_ssrf_guard(url):
    assert not web.is_public_url(url)


def test_fetcher_blocks_private_redirect(monkeypatch):
    def handler(req):
        return httpx.Response(302, headers={"location": "http://127.0.0.1/admin"})

    f = web.Fetcher(client=httpx.Client(transport=httpx.MockTransport(handler)), url_check=lambda u: "127.0.0.1" not in u)
    with pytest.raises(ValueError, match="blocked_url"):
        f.fetch_bytes("https://example.org/start")


# ---------------------------------------------------------------- harvesting
def _mock_http(monkeypatch, handler):
    real = httpx.Client

    def factory(*a, **k):
        k.pop("transport", None)
        return real(*a, transport=httpx.MockTransport(handler), **k)

    monkeypatch.setattr("app.plagiarism.harvest.base.httpx.Client", factory)
    monkeypatch.setattr(web, "is_public_url", lambda u: True)
    monkeypatch.setattr("app.plagiarism.web.httpx.Client", factory)


def _pdf_bytes(text: str) -> bytes:
    pdf = pymupdf.open()
    page = pdf.new_page()
    from tests.fixtures.documents import font_path

    fp = font_path()
    if fp:
        page.insert_font(fontname="dv", fontfile=fp)
    page.insert_textbox(pymupdf.Rect(40, 40, 560, 800), text, fontsize=9, fontname="dv" if fp else "helv")
    return pdf.tobytes()


def test_harvest_openalex_crossref_core_ojs_cyberleninka(admin, monkeypatch):
    monkeypatch.setattr(config.get_settings(), "CORE_API_KEY", SecretStr("core-key"))
    en = SAMPLES["en"]["human"]
    inv = {}
    for i, w in enumerate(en[0].split()):
        inv.setdefault(w, []).append(i)
    oai = """<?xml version="1.0"?><OAI-PMH xmlns="http://www.openarchives.org/OAI/2.0/"><ListRecords><record><header><identifier>oai:j:1</identifier></header>
      <metadata><oai_dc:dc xmlns:oai_dc="http://www.openarchives.org/OAI/2.0/oai_dc/" xmlns:dc="http://purl.org/dc/elements/1.1/">
      <dc:title>Mustaqil ta'lim</dc:title><dc:creator>Karimov A.</dc:creator><dc:date>2023-01-01</dc:date>
      <dc:description>Qisqa annotatsiya.</dc:description><dc:identifier>https://jurnal.uz/index.php/j/article/view/12</dc:identifier>
      </oai_dc:dc></metadata></record></ListRecords></OAI-PMH>"""

    def handler(req: httpx.Request):
        u = str(req.url)
        if "openalex" in u:
            return httpx.Response(200, json={"meta": {"next_cursor": None}, "results": [
                {"title": "OpenAlex paper", "abstract_inverted_index": inv, "publication_year": 2022, "doi": "https://doi.org/10.1/oa1",
                 "authorships": [{"author": {"display_name": "A. Author"}}], "best_oa_location": {"pdf_url": None, "landing_page_url": "https://oa.org/1"}}]})
        if "crossref" in u:
            return httpx.Response(200, json={"message": {"items": [
                {"title": ["Crossref paper"], "abstract": f"<jats:p>{en[1]}</jats:p>", "DOI": "10.1/cr1", "URL": "https://doi.org/10.1/cr1",
                 "issued": {"date-parts": [[2021]]}, "author": [{"given": "B", "family": "Writer"}]}]}} if "offset=0" in u else {"message": {"items": []}})
        if "core.ac.uk" in u:
            assert req.headers["authorization"] == "Bearer core-key"
            return httpx.Response(200, json={"results": [{"title": "CORE paper", "fullText": en[2] + " " + en[3], "yearPublished": 2020, "doi": "10.1/core1"}]} if "offset=0" in u else {"results": []})
        if u.endswith("/pdf"):
            return httpx.Response(200, headers={"content-type": "application/pdf"}, content=_pdf_bytes(SAMPLES["ru"]["human"][1] + " " + SAMPLES["ru"]["human"][2]))
        if "cyberleninka" in u:
            body = json.loads(req.content)
            arts = [{"name": "<b>Киберленинка</b> статья", "annotation": SAMPLES["ru"]["human"][0], "link": "/article/n/test", "authors": ["Иванов"], "year": 2019}] if body["from"] == 0 else []
            return httpx.Response(200, json={"articles": arts})
        if "/oai" in u:
            return httpx.Response(200, content=oai.encode())
        if "/article/view/12" in u:
            return httpx.Response(200, text='<a class="obj_galley_link pdf" href="https://jurnal.uz/index.php/j/article/view/12/34">PDF</a>')
        if "/article/download/12/34" in u:
            return httpx.Response(200, headers={"content-type": "application/pdf"}, content=_pdf_bytes(UZ["human"][0] + " " + UZ["human"][1]))
        return httpx.Response(404)

    _mock_http(monkeypatch, handler)
    r = admin.post("/api/corpus/harvest", json={"sources": ["openalex", "crossref", "core", "cyberleninka", "ojs"], "queries": ["digital learning"],
                                                "ojs_urls": ["https://jurnal.uz/index.php/j"], "limit": 5})
    assert r.status_code == 201, r.text
    job = admin.get(f"/api/corpus/jobs/{r.json()['id']}").json()
    assert job["status"] == "completed", job
    assert job["added"] == 5, job["log"]
    docs = admin.get("/api/corpus/documents", params={"size": 50}).json()["items"]
    by_src = {d["source_type"]: d for d in docs}
    assert set(by_src) == {"openalex", "crossref", "core", "cyberleninka", "ojs"}
    assert by_src["ojs"]["fulltext"] and by_src["cyberleninka"]["fulltext"] and by_src["core"]["fulltext"]
    assert not by_src["openalex"]["fulltext"]  # abstract only
    assert by_src["crossref"]["doi"] == "10.1/cr1"
    # harvesting again skips what is already there (DOI / URL / content)
    r2 = admin.post("/api/corpus/harvest", json={"sources": ["crossref", "core"], "queries": ["digital learning"], "limit": 5})
    job2 = admin.get(f"/api/corpus/jobs/{r2.json()['id']}").json()
    assert job2["added"] == 0 and job2["skipped"] == 2


def test_harvest_source_error_is_logged_not_fatal(admin, monkeypatch):
    _mock_http(monkeypatch, lambda req: httpx.Response(503))
    r = admin.post("/api/corpus/harvest", json={"sources": ["openalex"], "queries": ["x"], "limit": 3})
    job = admin.get(f"/api/corpus/jobs/{r.json()['id']}").json()
    assert job["status"] == "completed" and job["failed"] == 1 and any("openalex" in line for line in job["log"])


# ---------------------------------------------------------------- report
def test_plagiarism_pdf_report(admin):
    corpus_upload(admin, [("src/r.docx", make_docx(UZ["human"]))])
    aid = upload(admin, "report.docx", make_docx(["O'z matnim, yangi fikrlar bilan boyitilgan. " * 5, UZ["human"][1], UZ["human"][2]]))["analysis"]["id"]
    r = admin.get(f"/api/analyses/{aid}/report", params={"kind": "plagiarism"})
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf" and "antiplagiat" in r.headers["content-disposition"]
    text = "".join(p.get_text() for p in pymupdf.open(stream=r.content, filetype="pdf"))
    for needle in ("O'ZLASHTIRISHLARNI TEKSHIRISH HISOBOTI", "ORIGINALLIK", "O'ZLASHTIRISH", "IQTIBOSLAR", "MANBALAR", "TEKSHIRILGAN MATN", "Ma'lumotnoma bazasi"):
        assert needle in text, needle
    assert "[1]" in text  # source markers in the highlighted text


def test_corpus_document_with_single_title_heading_is_indexed(admin):
    """Regression (found in E2E): a file whose only heading is its title must still be indexed."""
    res = corpus_upload(admin, [("d/one_title.docx", make_docx(UZ["human"], {0: "Raqamli ta'lim asoslari"}))])
    job = admin.get(f"/api/corpus/jobs/{res['job']['id']}").json()
    assert job["added"] == 1, job["log"]


def test_brave_connection_error_is_explained_and_stops(user_client, monkeypatch):
    class Offline(FakeBrave):
        def search(self, q):
            Offline.calls += 1
            raise httpx.ConnectError("proxy refused")

    Offline.calls = 0
    monkeypatch.setattr(config.get_settings(), "BRAVE_API_KEY", SecretStr("test-key"))
    monkeypatch.setattr(web, "BraveClient", Offline)
    aid = upload(user_client, "offline.docx", make_docx(UZ["human"]), web_check="true", keep_for_similarity="false")["analysis"]["id"]
    user_client.post(f"/api/analyses/{aid}/confirm", json={"web_check": True})
    stats = user_client.get(f"/api/analyses/{aid}/plagiarism").json()["modules"]["web_stats"]
    assert Offline.calls == 1 and stats["queries_used"] == 0 and stats["errors"][0].startswith("brave_connection")
