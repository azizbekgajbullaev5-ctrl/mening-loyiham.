"""DOCX report (python-docx)."""
from __future__ import annotations

import io

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor

from app.reporting.builder import fmt_dt, fmt_pct

BLUE = RGBColor(0x1D, 0x4E, 0xD8)
GREY = RGBColor(0x64, 0x74, 0x8B)


def _table(d, rows, header=True):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = str(val)
            if header and i == 0:
                for r in cell.paragraphs[0].runs:
                    r.bold = True
    d.add_paragraph()
    return t


def _note(d, text):
    p = d.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.color.rgb = BLUE
    return p


def _small(d, text):
    p = d.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(8.5)
    run.font.color.rgb = GREY
    return p


def render_docx(data: dict) -> bytes:
    L = data["L"]
    r = data["detail"]["result"] or {}
    d = docx.Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10)
    d.add_heading(L["title"], level=0)
    _note(d, data["disclaimer"])

    doc = data["doc"]
    d.add_heading(L["s_doc"], level=1)
    _table(d, [
        [L["file"], doc["file"]], [L["doc_type"], doc["doc_type"]], [L["size"], f"{doc['size_kb']} KB ({doc['file_type']})"],
        [L["language"], doc["language"]], [L["words"], doc["words"]], [L["chars"], doc["chars"]],
        [L["pages"], f"{doc['pages']} {L['pages_est'] if doc['pages_estimated'] else ''}"], [L["depth"], data["depth"]],
        [L["analysis_date"], fmt_dt(data["analysis_date"])], [L["report_date"], fmt_dt(data["report_date"])],
    ], header=False)

    ai = r.get("ai", {})
    d.add_heading(L["s_ai"], level=1)
    _table(d, [
        [L["ai_overall"], L["confidence"], L["flagged"]],
        ["—" if ai.get("likelihood") is None else f"{ai['likelihood']:.0f}%", ai.get("confidence_label", "—"), f"{ai.get('passages_flagged', 0)} / {ai.get('passages_analyzed', 0)}"],
    ])
    _note(d, ai.get("note", ""))
    _small(d, f"{L['basis']}: {ai.get('basis_text', '')}")
    if ai.get("language_note"):
        _small(d, ai["language_note"])
    comp = r.get("provider_comparison", {})
    d.add_heading(L["s_providers"], level=2)
    rows = [["Provider", "Status", L["ai"]]]
    for p in r.get("providers", []):
        mean = ai.get("likelihood") if p.get("kind") == "local" else next((c.get("mean_score") for c in comp.get("providers", []) if c["name"] == p["name"]), None)
        rows.append([p["name"], p.get("status_label", p.get("status")), "—" if mean is None else f"{mean:.0f}%"])
    _table(d, rows)
    _small(d, comp.get("summary_text", ""))

    sim = r.get("similarity", {})
    d.add_heading(L["s_sim"], level=1)
    ext = fmt_pct(sim.get("external")) if sim.get("external") is not None else (L["not_configured"] if sim.get("scope") == "local_only" else L["not_run"])
    _table(d, [
        [L["sim_overall"], fmt_pct(sim.get("overall"))], [L["sim_internal"], fmt_pct(sim.get("internal"))],
        [L["sim_corpus"], fmt_pct(sim.get("corpus")) if sim.get("corpus") is not None else L["not_run"]], [L["sim_external"], ext],
    ], header=False)
    _note(d, sim.get("scope_text", ""))

    d.add_heading(L["s_chapters"], level=1)
    rows = [[L["section"], L["pages"], L["words"], L["ai"], L["confidence"], L["sim"]]]
    for s in data["sections"]:
        rows.append([
            "   " * (s["level"] - 1) + (s["title"] or s["kind_label"])[:90], f"{s['page_start'] or '—'}–{s['page_end'] or '—'}", s["word_count"],
            L["excluded"] if s["excluded_from_ai"] else fmt_pct(s["ai_likelihood"]), s["ai_confidence_label"] or "—", fmt_pct(s["similarity"]),
        ])
    _table(d, rows)

    d.add_heading(L["s_passages"], level=1)
    if not data["passages"]:
        d.add_paragraph(L["no_passages"])
    for p in data["passages"]:
        _small(d, f"{L['page']} {p['page'] or '—'} · {L['paragraph']} {p['paragraph_number']} · {p['section_title'] or ''}")
        para = d.add_paragraph()
        para.add_run(f"{L['ai']}: {p['ai_likelihood']:.0f}% · {L['confidence']}: {p['confidence_label']}").bold = True
        q = d.add_paragraph(p["text"])
        q.paragraph_format.left_indent = Pt(14)
        _small(d, f"{L['characteristics']}: " + ("; ".join(c["label"] for c in p["characteristics"]) or "—"))
        _small(d, p["explanation"])

    d.add_heading(L["s_matches"], level=1)
    if not data["matches"]:
        d.add_paragraph(L["no_matches"])
    for m in data["matches"]:
        src = m["source_title"] or m["matched_document_name"] or (f"{L['page']} {m['matched_page']}, {L['paragraph']} {m['matched_paragraph_number']}" if m["matched_paragraph_number"] else "")
        _small(d, f"{m['match_type_label']} — {m['similarity']:.0f}% · {L['page']} {m['page'] or '—'}, {L['paragraph']} {m['paragraph_number'] or '—'} ↔ {src}")
        d.add_paragraph(m["text"][:500] + ("…" if len(m["text"]) > 500 else ""))
        if m["source_url"]:
            _small(d, m["source_url"])
    rp = r.get("repeated_phrases") or []
    if rp:
        d.add_heading(L["repeated"], level=2)
        _table(d, [["", "×"]] + [[x["phrase"], x["count"]] for x in rp[:15]])

    ac = r.get("academic", {})
    d.add_heading(L["s_academic"], level=1)
    if ac.get("metrics"):
        _small(d, " · ".join(f"{k}: {v}" for k, v in ac["metrics"].items()))
    issues = ac.get("issues", [])
    if issues:
        _table(d, [["", ""]] + [[i["severity_label"], i["text"]] for i in issues])
    else:
        d.add_paragraph(L["no_issues"])

    d.add_heading(L["s_method"], level=1)
    for x in data["methodology"]:
        d.add_paragraph(x)
    d.add_heading(L["s_limits"], level=1)
    for x in data["limitations"]:
        d.add_paragraph(x, style="List Bullet")
    d.add_heading(L["s_disclaimer"], level=1)
    _note(d, data["disclaimer"])

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
