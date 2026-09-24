"""Text extraction for DOCX, PDF (with OCR fallback) and TXT."""
from __future__ import annotations

import io
import logging
import re
import statistics
import zipfile
from collections import Counter
from collections.abc import Callable

from app.core.config import get_settings
from app.document_processing.types import WORDS_PER_PAGE_ESTIMATE, Block, ExtractedDocument, ExtractionError
from app.analyzers.text_utils import collapse_ws, words

log = logging.getLogger(__name__)
ProgressFn = Callable[[float, str], None]

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


OcrPageFn = Callable[[int, str], None]


def extract(
    file_type: str,
    data: bytes,
    progress: ProgressFn | None = None,
    ocr_cache: dict[str, str] | None = None,
    on_ocr_page: OcrPageFn | None = None,
) -> ExtractedDocument:
    """``ocr_cache`` / ``on_ocr_page`` let a resumed analysis skip pages already OCR'd."""
    if file_type == "docx":
        if progress:
            progress(0.1, "docx")
        doc = extract_docx(data)
    elif file_type == "pdf":
        doc = extract_pdf(data, progress, ocr_cache, on_ocr_page)
    elif file_type == "txt":
        doc = extract_txt(data)
    else:
        raise ExtractionError("unsupported_type", f"Unsupported type {file_type}")
    for i, b in enumerate(doc.blocks):
        b.index = i
    if not any(words(b.text) for b in doc.blocks):
        if doc.is_scanned and not doc.ocr_used:
            raise ExtractionError("scanned_no_ocr", "Scanned PDF detected but OCR is unavailable")
        raise ExtractionError("no_text", "No extractable text found in the document")
    return doc


# ---------------------------------------------------------------- DOCX
def _heading_level_from_style(style_name: str | None) -> int | None:
    if not style_name:
        return None
    s = style_name.lower()
    if s in {"title", "название", "sarlavha nomi"}:
        return 0
    m = re.match(r"(heading|заголовок|sarlavha)\s*(\d)", s)
    if m:
        return int(m.group(2))
    return None


def extract_docx(data: bytes) -> ExtractedDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - library raises many types
        raise ExtractionError("parse_error", f"Could not parse DOCX: {exc}") from exc

    blocks: list[Block] = []
    hidden: list[dict] = []
    page = 1
    saw_page_marker = False
    for child in document.element.body.iterchildren():
        tag = child.tag
        if tag == W_NS + "p":
            para = Paragraph(child, document)
            ppr = child.find(W_NS + "pPr")
            before_text, after_text = _count_page_breaks(child)
            if ppr is not None and ppr.find(W_NS + "pageBreakBefore") is not None:
                before_text = max(before_text, 1)
            if before_text or after_text:
                saw_page_marker = True
            page += before_text
            text = collapse_ws(_visible_paragraph_text(para, hidden, page))
            page_after = page + after_text
            if not text:
                page = page_after
                continue
            style = para.style.name if para.style is not None else None
            level = _heading_level_from_style(style)
            if level is None and ppr is not None:
                ol = ppr.find(W_NS + "outlineLvl")
                if ol is not None:
                    try:
                        level = int(ol.get(W_NS + "val")) + 1
                    except (TypeError, ValueError):
                        level = None
            kind = "heading" if level is not None else ("list" if style and "list" in style.lower() else "paragraph")
            blocks.append(Block(index=len(blocks), text=text, kind=kind, page=page, heading_level=level, style=style))
            page = page_after
            if ppr is not None and ppr.find(f"{W_NS}sectPr") is not None:
                sect_type = ppr.find(f"{W_NS}sectPr/{W_NS}type")
                if sect_type is None or sect_type.get(W_NS + "val") in (None, "nextPage", "oddPage", "evenPage"):
                    page += 1
                    saw_page_marker = True
        elif tag == W_NS + "tbl":
            table = Table(child, document)
            rows = []
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    t = collapse_ws(cell.text)
                    if t and (not cells or cells[-1] != t):  # merged cells repeat
                        cells.append(t)
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                blocks.append(Block(index=len(blocks), text="\n".join(rows), kind="table", page=page))

    declared_pages = _docx_declared_pages(data)
    warnings: list[str] = []
    total_words = sum(len(words(b.text)) for b in blocks)
    if saw_page_marker:
        page_count = max(page, declared_pages or 0)
        pages_estimated = False
    else:
        estimate = max(1, round(total_words / WORDS_PER_PAGE_ESTIMATE))
        # docProps "Pages" is only refreshed when Word saves the file; generated files often carry a stale value.
        plausible = declared_pages and estimate / 3 <= declared_pages <= estimate * 3
        page_count = declared_pages if plausible else estimate
        pages_estimated = True
        _assign_estimated_pages(blocks, page_count)
        warnings.append("docx_pages_estimated")
    return ExtractedDocument("docx", blocks, page_count, pages_estimated=pages_estimated, warnings=warnings, hidden_fragments=hidden[:500])


_WHITE = {"FFFFFF", "FFFFFE", "FEFEFE", "FDFDFD"}


def _visible_paragraph_text(para, hidden: list[dict], page: int) -> str:
    """Paragraph text without runs a reader cannot see (white, hidden, <=2pt)."""
    from docx.text.run import Run

    parts: list[str] = []
    removed = False
    runs = [Run(r, para) for r in para._p.iter(W_NS + "r") if not any(a.tag in (W_NS + "del", W_NS + "moveFrom") for a in r.iterancestors())]
    for run in runs:
        t = run.text
        if not t:
            continue
        reason = None
        rpr = run._element.rPr
        if rpr is not None:
            if rpr.find(W_NS + "vanish") is not None or rpr.find(W_NS + "specVanish") is not None:
                reason = "hidden"
            col = rpr.find(W_NS + "color")
            if col is not None and (col.get(W_NS + "val") or "").upper() in _WHITE:
                reason = reason or "white"
            sz = rpr.find(W_NS + "sz")
            try:
                if sz is not None and int(sz.get(W_NS + "val")) <= 4:  # half-points: <= 2pt
                    reason = reason or "tiny"
            except (TypeError, ValueError):
                pass
        if reason and t.strip():
            hidden.append({"page": page, "reason": reason, "chars": len(t), "sample": t.strip()[:80]})
            removed = True
            parts.append(" ")
            continue
        parts.append(t)
    if not runs:
        return para.text
    joined = "".join(parts)
    # hyperlinks / fields are not in para.runs: fall back to full text when runs miss content
    if removed or joined.strip():
        return joined
    return para.text


def _count_page_breaks(p_elem) -> tuple[int, int]:
    """Page breaks before the paragraph's first text vs. after it.

    Word stores both hard breaks (<w:br w:type="page">) and, in saved files,
    <w:lastRenderedPageBreak/> markers where its layout broke the page. A
    rendered break adjacent to a hard break describes the same break.
    """
    before = after = 0
    seen_text = False
    last_was_break = False
    for el in p_elem.iter():
        if el.tag == W_NS + "t" and el.text and el.text.strip():
            seen_text = True
            last_was_break = False
        elif (el.tag == W_NS + "br" and el.get(W_NS + "type") == "page") or el.tag == W_NS + "lastRenderedPageBreak":
            if last_was_break:
                continue
            last_was_break = True
            if seen_text:
                after += 1
            else:
                before += 1
    return before, after


def _docx_declared_pages(data: bytes) -> int | None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            if "docProps/app.xml" not in zf.namelist():
                return None
            m = re.search(rb"<Pages>(\d+)</Pages>", zf.read("docProps/app.xml"))
            return int(m.group(1)) if m else None
    except (zipfile.BadZipFile, KeyError):
        return None


def _assign_estimated_pages(blocks: list[Block], page_count: int) -> None:
    total = sum(max(1, len(words(b.text))) for b in blocks) or 1
    acc = 0
    for b in blocks:
        b.page = min(page_count, 1 + int(acc / total * page_count))
        acc += max(1, len(words(b.text)))


# ---------------------------------------------------------------- PDF
def extract_pdf(
    data: bytes,
    progress: ProgressFn | None = None,
    ocr_cache: dict[str, str] | None = None,
    on_ocr_page: OcrPageFn | None = None,
) -> ExtractedDocument:
    import pymupdf as fitz

    try:
        pdf = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("parse_error", f"Could not parse PDF: {exc}") from exc
    if pdf.needs_pass:
        raise ExtractionError("encrypted", "Password-protected PDF files are not supported")

    s = get_settings()
    page_count = pdf.page_count
    raw_pages: list[list[dict]] = []
    scanned_pages: list[int] = []
    hidden: list[dict] = []
    for pno in range(page_count):
        page = pdf[pno]
        height = page.rect.height or 1
        items = []
        for blk in page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE)["blocks"]:
            if blk.get("type") != 0:
                continue
            lines, sizes, bold_chars, chars = [], [], 0, 0
            for line in blk["lines"]:
                visible = []
                for sp in line["spans"]:
                    if sp["text"].strip() and (sp.get("color") in _PDF_WHITE or sp["size"] < 2.0):
                        hidden.append({"page": pno + 1, "reason": "white" if sp.get("color") in _PDF_WHITE else "tiny",
                                       "chars": len(sp["text"]), "sample": sp["text"].strip()[:80]})
                        continue
                    visible.append(sp)
                spans = [sp for sp in visible if sp["text"].strip()]
                if not spans:
                    continue
                lines.append("".join(sp["text"] for sp in visible))
                for sp in spans:
                    n = len(sp["text"].strip())
                    sizes.append((sp["size"], n))
                    chars += n
                    if sp["flags"] & 16:
                        bold_chars += n
            if not lines:
                continue
            text = _join_pdf_lines(lines)
            size = max(sizes, key=lambda x: x[1])[0] if sizes else 0
            y0, y1 = blk["bbox"][1], blk["bbox"][3]
            items.append(
                {
                    "text": text,
                    "size": size,
                    "bold": chars > 0 and bold_chars / chars > 0.8,
                    "edge": y1 < height * 0.08 or y0 > height * 0.92,
                    "chars": chars,
                }
            )
        page_chars = sum(i["chars"] for i in items)
        if page_chars < 25 and page.get_images():
            scanned_pages.append(pno)
        raw_pages.append(items)
        if progress and pno % 10 == 0:
            progress(pno / max(1, page_count) * 0.5, f"{pno + 1}/{page_count}")

    warnings: list[str] = []
    is_scanned = bool(page_count) and len(scanned_pages) / page_count >= 0.5
    ocr_used = False
    if scanned_pages:
        if s.OCR_ENABLED and _ocr_available():
            todo = scanned_pages[: s.OCR_MAX_PAGES]
            if len(todo) < len(scanned_pages):
                warnings.append("ocr_page_limit")
            ocr_cache = ocr_cache or {}
            for k, pno in enumerate(todo):
                if str(pno) in ocr_cache:  # resumed: this page was already recognised
                    text = ocr_cache[str(pno)]
                else:
                    text = _ocr_page(pdf[pno], s.OCR_DPI, s.OCR_LANGUAGES)
                    if on_ocr_page:
                        on_ocr_page(pno, text)
                raw_pages[pno] = [
                    {"text": collapse_ws(p), "size": 0, "bold": False, "edge": False, "chars": len(p), "ocr": True}
                    for p in re.split(r"\n\s*\n", text)
                    if p.strip()
                ]
                if progress:
                    progress(0.5 + (k + 1) / len(todo) * 0.5, f"OCR {k + 1}/{len(todo)}")
            ocr_used = True
            warnings.append("ocr_used")
        else:
            warnings.append("scanned_pages_without_ocr")

    repeated = _repeated_edge_lines(raw_pages)
    body_sizes = [i["size"] for items in raw_pages for i in items for _ in range(max(1, i["chars"] // 50)) if i["size"]]
    body = statistics.median(body_sizes) if body_sizes else 0
    distinct_big = sorted({round(i["size"]) for items in raw_pages for i in items if body and i["size"] >= body * 1.15}, reverse=True)

    blocks: list[Block] = []
    for pno, items in enumerate(raw_pages, start=1):
        for it in items:
            text = it["text"]
            if not text:
                continue
            if it["edge"] and (re.fullmatch(r"[\d\s\-–—|/.]+", text) or _edge_key(text) in repeated):
                continue  # running header/footer or page number
            level = None
            short = len(text) <= 160 and not text.rstrip().endswith((".", ",", ";"))
            if body and it["size"] >= body * 1.15 and short:
                level = min(3, distinct_big.index(round(it["size"])) + 1) if round(it["size"]) in distinct_big else 1
            elif it["bold"] and short and len(words(text)) <= 20:
                level = 3
            blocks.append(
                Block(
                    index=len(blocks),
                    text=text,
                    kind="heading" if level else "paragraph",
                    page=pno,
                    heading_level=level,
                    is_ocr=bool(it.get("ocr")),
                )
            )
    blocks = _merge_split_pdf_paragraphs(blocks)
    return ExtractedDocument(
        "pdf", blocks, page_count, pages_estimated=False, is_scanned=is_scanned, ocr_used=ocr_used, warnings=warnings,
        hidden_fragments=hidden[:500],
    )


_PDF_WHITE = {0xFFFFFF, 0xFEFEFE, 0xFDFDFD}


def _join_pdf_lines(lines: list[str]) -> str:
    out = ""
    for ln in lines:
        ln = ln.strip()
        if not ln:
            continue
        if out.endswith("-") and ln[:1].islower():
            out = out[:-1] + ln  # de-hyphenate
        else:
            out = (out + " " + ln) if out else ln
    return collapse_ws(out)


def _edge_key(text: str) -> str:
    return re.sub(r"\d+", "#", text.lower()).strip()


def _repeated_edge_lines(pages: list[list[dict]]) -> set[str]:
    if len(pages) < 3:
        return set()
    counts = Counter()
    for items in pages:
        counts.update({_edge_key(i["text"]) for i in items if i["edge"]})
    return {k for k, c in counts.items() if c >= max(3, len(pages) * 0.3)}


def _merge_split_pdf_paragraphs(blocks: list[Block]) -> list[Block]:
    """Join a paragraph that continues across a page boundary."""
    out: list[Block] = []
    for b in blocks:
        prev = out[-1] if out else None
        if (
            prev
            and prev.kind == "paragraph"
            and b.kind == "paragraph"
            and prev.page is not None
            and b.page == prev.page + 1
            and not re.search(r"[.!?:»\"”)]$", prev.text)
            and b.text[:1].islower()
        ):
            prev.text = _join_pdf_lines([prev.text, b.text])
            continue
        out.append(b)
    for i, b in enumerate(out):
        b.index = i
    return out


_WINDOWS_TESSERACT = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
)
_ocr_langs: set[str] | None = None


def _configure_tesseract() -> None:
    import os

    import pytesseract

    cmd = get_settings().TESSERACT_CMD
    if not cmd and os.name == "nt":
        local = os.path.expandvars(r"%LOCALAPPDATA%\Programs\Tesseract-OCR\tesseract.exe")
        cmd = next((c for c in (*_WINDOWS_TESSERACT, local) if os.path.exists(c)), "")
    if cmd:
        pytesseract.pytesseract.tesseract_cmd = cmd


def _ocr_available() -> bool:
    try:
        import pytesseract

        _configure_tesseract()
        pytesseract.get_tesseract_version()
        return True
    except Exception:  # noqa: BLE001
        return False


def _ocr_page(page, dpi: int, langs: str) -> str:
    import pytesseract
    from PIL import Image

    pix = page.get_pixmap(dpi=dpi)
    img = Image.open(io.BytesIO(pix.tobytes("png")))
    global _ocr_langs
    if _ocr_langs is None:
        _ocr_langs = set(pytesseract.get_languages(config=""))
    available = _ocr_langs
    use = "+".join(lang for lang in langs.split("+") if lang in available) or "eng"
    return pytesseract.image_to_string(img, lang=use)


# ---------------------------------------------------------------- TXT
def decode_text(data: bytes) -> str:
    for bom, enc in ((b"\xef\xbb\xbf", "utf-8-sig"), (b"\xff\xfe", "utf-16"), (b"\xfe\xff", "utf-16")):
        if data.startswith(bom):
            return data.decode(enc)
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        pass
    try:
        text = data.decode("cp1251")  # common for legacy Russian/Uzbek files
        if re.search(r"[а-яА-Я]", text):
            return text
    except UnicodeDecodeError:
        pass
    return data.decode("latin-1")


def extract_txt(data: bytes) -> ExtractedDocument:
    text = decode_text(data).replace("\r\n", "\n").replace("\r", "\n")
    pages = text.split("\f")
    has_ff = len(pages) > 1
    blocks: list[Block] = []
    for pno, ptext in enumerate(pages, start=1):
        lines = ptext.split("\n")
        non_empty = [ln for ln in lines if ln.strip()]
        blank_separated = ptext.count("\n\n") >= max(1, len(non_empty) // 10)
        paras = re.split(r"\n\s*\n", ptext) if blank_separated else non_empty
        for p in paras:
            t = collapse_ws(p)
            if t:
                blocks.append(Block(index=len(blocks), text=t, page=pno if has_ff else None))
    if has_ff:
        page_count, estimated = len(pages), False
    else:
        total_words = sum(len(words(b.text)) for b in blocks)
        page_count = max(1, round(total_words / WORDS_PER_PAGE_ESTIMATE))
        estimated = True
        _assign_estimated_pages(blocks, page_count)
    return ExtractedDocument("txt", blocks, page_count, pages_estimated=estimated, warnings=["txt_pages_estimated"] if estimated else [])
